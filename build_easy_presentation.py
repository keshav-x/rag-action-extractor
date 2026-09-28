import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Color Palette (Prismatic Editorial Theme)
    BG_CREAM = RGBColor(250, 248, 245)
    TEXT_TITLE = RGBColor(30, 27, 38)        # Deep Slate/Plum
    TEXT_BODY = RGBColor(55, 53, 64)         # Charcoal
    TEXT_MUTED = RGBColor(112, 108, 120)     # Muted Eyebrow
    CARD_BG = RGBColor(255, 255, 255)        # Pure White Card
    CARD_BORDER = RGBColor(226, 222, 215)    # Subtle warm gray
    ACCENT_CORAL = RGBColor(215, 65, 85)     # Editorial Coral
    ACCENT_GREEN = RGBColor(35, 134, 54)     # Success Green
    ACCENT_SLATE = RGBColor(45, 55, 72)      # Tech Slate

    bg_image_path = os.path.abspath("prismatic_bg.jpg")
    screenshot_main = os.path.abspath("slide_screenshot_main.png")
    screenshot_indexed = os.path.abspath("slide_screenshot_indexed.png")
    screenshot_results = os.path.abspath("slide_screenshot_results.png")

    def add_bg(slide, use_prismatic=False):
        if use_prismatic and os.path.exists(bg_image_path):
            slide.shapes.add_picture(bg_image_path, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
        else:
            bg_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
            bg_rect.fill.solid()
            bg_rect.fill.fore_color.rgb = BG_CREAM
            bg_rect.line.fill.background()

    def add_header(slide, eyebrow, title):
        # Eyebrow
        tx_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.35))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = eyebrow.upper()
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CORAL

        # Title
        tx_title = slide.shapes.add_textbox(Inches(0.9), Inches(0.9), Inches(11.5), Inches(0.75))
        tf_title = tx_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_t = tf_title.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Georgia"
        p_t.font.size = Pt(27)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_TITLE

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, use_prismatic=True)

    # Eyebrow tag
    tb_tag = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(8.0), Inches(0.4))
    tf_tag = tb_tag.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "ACADEMIC CAPSTONE PROJECT • CLASS DEMO"
    p_tag.font.name = "Segoe UI"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CORAL

    # Main Title
    tb_title = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(8.5), Inches(2.2))
    tft = tb_title.text_frame
    tft.word_wrap = True
    p_title = tft.paragraphs[0]
    p_title.text = "Document Action\nExtractor"
    p_title.font.name = "Georgia"
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_TITLE

    # Subtitle
    tb_sub = s1.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(8.0), Inches(0.8))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Autonomous Task, Assignee & Deadline Extraction from Unstructured Text\nPowered by Retrieval-Augmented Generation (RAG) & Local LLM Inference"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_BODY

    # Team Members Card
    team_card = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(5.5), Inches(7.5), Inches(1.3))
    team_card.fill.solid()
    team_card.fill.fore_color.rgb = CARD_BG
    team_card.line.color.rgb = CARD_BORDER
    team_card.line.width = Pt(1)

    tftc = team_card.text_frame
    tftc.margin_left = tftc.margin_top = tftc.margin_right = tftc.margin_bottom = Inches(0.2)
    tftc.word_wrap = True

    pt1 = tftc.paragraphs[0]
    pt1.text = "PROJECT AUTHORS & CONTRIBUTORS"
    pt1.font.name = "Segoe UI"
    pt1.font.size = Pt(10)
    pt1.font.bold = True
    pt1.font.color.rgb = ACCENT_CORAL

    pt2 = tftc.add_paragraph()
    pt2.text = "• Keshav — Roll No: 2449391        • Darshveer — Roll No: 2449377"
    pt2.font.name = "Segoe UI"
    pt2.font.size = Pt(13)
    pt2.font.bold = True
    pt2.font.color.rgb = TEXT_TITLE
    pt2.space_before = Pt(4)

    pt3 = tftc.add_paragraph()
    pt3.text = "Department of Computer Science & Engineering | Academic Submission"
    pt3.font.name = "Segoe UI"
    pt3.font.size = Pt(10.5)
    pt3.font.color.rgb = TEXT_MUTED
    pt3.space_before = Pt(2)

    # =========================================================================
    # SLIDE 2: THE PROBLEM (WHY THIS PROJECT MATTERS)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2, use_prismatic=False)
    add_header(s2, "01 / Motivation & Problem", "The Challenge of Unstructured Business & Technical Documents")

    problems = [
        ("01", "Information Overload", "Meeting minutes, technical specifications, and project updates span dozens of pages. Concrete tasks and critical decisions get lost in long conversational narratives."),
        ("02", "Diffused Ownership & Deadlines", "Action items mentioned informally ('we should test this by Friday') often lack explicit assignee tracking, leading to forgotten deadlines and milestone delays."),
        ("03", "Manual Inefficiency", "Teams waste hours every week reading transcripts to manually copy tasks into Jira or Trello, a slow process heavily prone to human oversight.")
    ]

    for i, (num, p_title, p_desc) in enumerate(problems):
        card = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9 + i * 3.9), Inches(2.0), Inches(3.7), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tfc = card.text_frame
        tfc.margin_left = tfc.margin_top = tfc.margin_right = tfc.margin_bottom = Inches(0.3)
        tfc.word_wrap = True

        p_num = tfc.paragraphs[0]
        p_num.text = num
        p_num.font.name = "Georgia"
        p_num.font.size = Pt(28)
        p_num.font.bold = True
        p_num.font.color.rgb = ACCENT_CORAL

        p_head = tfc.add_paragraph()
        p_head.text = p_title
        p_head.font.name = "Georgia"
        p_head.font.size = Pt(17)
        p_head.font.bold = True
        p_head.font.color.rgb = TEXT_TITLE
        p_head.space_before = Pt(8)

        p_body = tfc.add_paragraph()
        p_body.text = p_desc
        p_body.font.name = "Segoe UI"
        p_body.font.size = Pt(12.5)
        p_body.font.color.rgb = TEXT_BODY
        p_body.space_before = Pt(12)

    # =========================================================================
    # SLIDE 3: THE SOLUTION & KEY GOALS
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3, use_prismatic=False)
    add_header(s3, "02 / The Solution", "Automated, Grounded Action Item Extraction with Zero Guesswork")

    # Left: Core Pillars Card
    left_card = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.9), Inches(5.6), Inches(4.9))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = CARD_BG
    left_card.line.color.rgb = CARD_BORDER
    left_card.line.width = Pt(1)

    tfl = left_card.text_frame
    tfl.margin_left = tfl.margin_top = tfl.margin_right = tfl.margin_bottom = Inches(0.3)
    tfl.word_wrap = True

    p = tfl.paragraphs[0]
    p.text = "What Does Document Action Extractor Do?"
    p.font.name = "Georgia"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    pillars = [
        ("Multi-Format Ingestion", "Uploads PDF, Word (.docx), and Text (.txt) files with smart document structure parsing."),
        ("Semantic Chunk Search", "Finds the exact sections discussing deliverables rather than passing noisy full text."),
        ("Structured JSON Output", "Extracts validated attributes: Task Description, Owner, Due Date, and Priority."),
        ("100% Auditable Source Citations", "Every single task card cites the exact sentence quote from the document.")
    ]

    for title, desc in pillars:
        p1 = tfl.add_paragraph()
        p1.text = "• " + title + ":"
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CORAL
        p1.space_before = Pt(10)

        p2 = tfl.add_paragraph()
        p2.text = "  " + desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(2)

    # Right: 3 Value Proposition Cards
    val_cards = [
        ("Privacy-First Local Execution", "Complete 100% offline capability using Ollama (Qwen 3.5 9B) and local ONNX embeddings. No confidential files ever leave your machine."),
        ("Instant Cloud Scale (Gemini)", "Seamless one-click toggle to Google Gemini Flash for lightning-fast (1.5s) cloud processing when offline privacy isn't required."),
        ("One-Click Universal Export", "Instantly download extracted tasks as CSV, formatted JSON, or Markdown for direct import into Notion, Jira, or GitHub Issues.")
    ]

    for i, (v_title, v_desc) in enumerate(val_cards):
        rcard = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.9 + i * 1.68), Inches(5.6), Inches(1.5))
        rcard.fill.solid()
        rcard.fill.fore_color.rgb = CARD_BG
        rcard.line.color.rgb = CARD_BORDER
        rcard.line.width = Pt(1)

        tfr = rcard.text_frame
        tfr.margin_left = tfr.margin_top = tfr.margin_right = tfr.margin_bottom = Inches(0.25)
        tfr.word_wrap = True

        pr1 = tfr.paragraphs[0]
        pr1.text = f"Feature 0{i+1}: {v_title}"
        pr1.font.name = "Georgia"
        pr1.font.size = Pt(14)
        pr1.font.bold = True
        pr1.font.color.rgb = TEXT_TITLE

        pr2 = tfr.add_paragraph()
        pr2.text = v_desc
        pr2.font.name = "Segoe UI"
        pr2.font.size = Pt(11.5)
        pr2.font.color.rgb = TEXT_BODY
        pr2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 4: HOW IT WORKS — THE 4-STAGE PIPELINE
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4, use_prismatic=False)
    add_header(s4, "03 / System Pipeline", "How Unstructured Text Becomes Structured Tasks")

    stages = [
        ("STAGE 1", "Document Ingestion", "PDF, Word, or TXT", "Document loaders extract raw text, and RecursiveCharacterTextSplitter segments prose into 800-character chunks with 100-character overlap."),
        ("STAGE 2", "Vector Indexing", "Local ONNX Embeddings", "The all-MiniLM-L6-v2 embedding model encodes each chunk into a 384-dimensional dense vector and stores it in persistent ChromaDB."),
        ("STAGE 3", "Similarity Retrieval", "ChromaDB Top-K Search", "User queries ('Extract all action items') are matched against stored vectors using cosine similarity, retrieving only the most relevant chunks."),
        ("STAGE 4", "LLM Task Extraction", "Pydantic JSON Extraction", "The retrieved context is passed to the LLM (Ollama / Gemini) with strict JSON constraints, generating validated task objects with verified quotes.")
    ]

    for i, (stg_tag, stg_title, stg_sub, stg_desc) in enumerate(stages):
        card = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9 + i * 2.92), Inches(2.0), Inches(2.75), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tfc = card.text_frame
        tfc.margin_left = tfc.margin_top = tfc.margin_right = tfc.margin_bottom = Inches(0.25)
        tfc.word_wrap = True

        p_tag = tfc.paragraphs[0]
        p_tag.text = stg_tag
        p_tag.font.name = "Segoe UI"
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_CORAL

        p_t = tfc.add_paragraph()
        p_t.text = stg_title
        p_t.font.name = "Georgia"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_TITLE
        p_t.space_before = Pt(4)

        p_sub = tfc.add_paragraph()
        p_sub.text = stg_sub
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = ACCENT_SLATE
        p_sub.space_before = Pt(2)

        p_desc = tfc.add_paragraph()
        p_desc.text = stg_desc
        p_desc.font.name = "Segoe UI"
        p_desc.font.size = Pt(11.5)
        p_desc.font.color.rgb = TEXT_BODY
        p_desc.space_before = Pt(12)

    # =========================================================================
    # SLIDE 5: SYSTEM ARCHITECTURE & TECH STACK
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5, use_prismatic=False)
    add_header(s5, "04 / Architecture & Tech Stack", "Robust Open-Source Technologies Built for Speed & Reliability")

    quads = [
        ("Presentation Layer", "Streamlit & Dark Slate CSS", [
            "Modern, distraction-free Dark Slate theme (#0e1117 / #161b22).",
            "Zero promotional clutter; high-contrast readability.",
            "Intuitive workflow: Document -> Query -> Visual Task Cards.",
            "Export options for CSV, JSON, and Markdown summaries."
        ]),
        ("Vector Database Engine", "ChromaDB Persistent Store", [
            "Embedded vector database requiring zero external cloud setup.",
            "Stores document embeddings with persistent disk storage.",
            "Sub-millisecond cosine distance similarity matching.",
            "Isolated collections allow re-indexing without data collisions."
        ]),
        ("Embedding Pipeline", "Local ONNX (all-MiniLM-L6-v2)", [
            "Runs 100% locally on CPU/GPU via ONNX runtime.",
            "Zero API calls, zero rate-limit errors (no 429 quota failures).",
            "Embeds 25 chunks in under 300ms.",
            "Cloud Google text-embedding-004 available as fallback."
        ]),
        ("Dual LLM Reasoning", "Ollama (Local) & Gemini (Cloud)", [
            "Ollama (Qwen 3.5 9B): 100% offline, privacy-guaranteed inference.",
            "Optimized with think=False to prevent empty-token hangs.",
            "Google Gemini Flash: High-speed cloud processing (~1.5s).",
            "Pydantic schema validation guarantees pure, structured JSON."
        ])
    ]

    for i, (q_cat, q_tech, q_bullets) in enumerate(quads):
        x = Inches(0.9 + (i % 2) * 5.9)
        y = Inches(1.9 + (i // 2) * 2.5)

        card = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(5.6), Inches(2.25))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tfc = card.text_frame
        tfc.margin_left = tfc.margin_top = tfc.margin_right = tfc.margin_bottom = Inches(0.2)
        tfc.word_wrap = True

        p_cat = tfc.paragraphs[0]
        p_cat.text = q_cat.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CORAL

        p_tech = tfc.add_paragraph()
        p_tech.text = q_tech
        p_tech.font.name = "Georgia"
        p_tech.font.size = Pt(15)
        p_tech.font.bold = True
        p_tech.font.color.rgb = TEXT_TITLE
        p_tech.space_before = Pt(2)

        for b in q_bullets:
            pb = tfc.add_paragraph()
            pb.text = "• " + b
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10.5)
            pb.font.color.rgb = TEXT_BODY
            pb.space_before = Pt(3)

    # =========================================================================
    # SLIDE 6: UI WALKTHROUGH — WORKSPACE & INGESTION
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6, use_prismatic=False)
    add_header(s6, "05 / User Interface", "Clean Ingestion Workspace & Real-Time Document Indexing")

    # Left: Screenshot
    if os.path.exists(screenshot_indexed):
        s6.shapes.add_picture(screenshot_indexed, Inches(0.9), Inches(1.9), width=Inches(6.8))

    # Right: Annotations Card
    rcard = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.0), Inches(1.9), Inches(4.4), Inches(4.9))
    rcard.fill.solid()
    rcard.fill.fore_color.rgb = CARD_BG
    rcard.line.color.rgb = CARD_BORDER
    rcard.line.width = Pt(1)

    tfr = rcard.text_frame
    tfr.margin_left = tfr.margin_top = tfr.margin_right = tfr.margin_bottom = Inches(0.3)
    tfr.word_wrap = True

    p = tfr.paragraphs[0]
    p.text = "Key Workspace Elements"
    p.font.name = "Georgia"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    annotations = [
        ("Sidebar Model Selector", "Switch easily between Local Ollama and Cloud Gemini. Choose temperature and top-K retrieved chunks."),
        ("Drag-and-Drop Ingestion", "Upload any PDF, DOCX, or TXT file. Automatic background chunking and vector storage with zero wait."),
        ("Document Text Preview", "An expandable inspector lets users verify the exact extracted text from multi-page documents."),
        ("Targeted Preset Queries", "One-click filter presets: 'All Action Items', 'High Priority Only', 'Upcoming Deadlines', or 'Assigned Tasks'.")
    ]

    for title, desc in annotations:
        p1 = tfr.add_paragraph()
        p1.text = "• " + title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CORAL
        p1.space_before = Pt(8)

        p2 = tfr.add_paragraph()
        p2.text = desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 7: UI WALKTHROUGH — STRUCTURED EXTRACTION IN ACTION
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7, use_prismatic=False)
    add_header(s7, "06 / Live Results", "Grounded Task Extraction with Verified Context Quotes")

    # Left: Results Screenshot
    if os.path.exists(screenshot_results):
        s7.shapes.add_picture(screenshot_results, Inches(0.9), Inches(1.9), width=Inches(6.8))

    # Right: Task Breakdown Card
    rcard7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.0), Inches(1.9), Inches(4.4), Inches(4.9))
    rcard7.fill.solid()
    rcard7.fill.fore_color.rgb = CARD_BG
    rcard7.line.color.rgb = CARD_BORDER
    rcard7.line.width = Pt(1)

    tfr7 = rcard7.text_frame
    tfr7.margin_left = tfr7.margin_top = tfr7.margin_right = tfr7.margin_bottom = Inches(0.3)
    tfr7.word_wrap = True

    p = tfr7.paragraphs[0]
    p.text = "Extracted Action Cards"
    p.font.name = "Georgia"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    results_notes = [
        ("Concrete Tasks Identified", "Extracted 4 distinct deliverables from sample.pdf (Alice, Bob, Charlie, and Diana) in ~7.5 seconds."),
        ("Explicit Metadata Mapping", "Each card clearly isolates the Responsible Owner, Due Date / Timeframe, and Priority Level badge."),
        ("Direct Text Verification", "The Source Citation box displays the exact original sentence from the document, eliminating hallucinations."),
        ("Multi-Format Export", "Buttons at the top allow instant export of all extracted items into CSV, JSON, or clean Markdown tables.")
    ]

    for title, desc in results_notes:
        p1 = tfr7.add_paragraph()
        p1.text = "• " + title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CORAL
        p1.space_before = Pt(8)

        p2 = tfr7.add_paragraph()
        p2.text = desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 8: ENGINEERING CHALLENGES & SOLUTIONS
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8, use_prismatic=False)
    add_header(s8, "07 / Technical Innovations", "Key Engineering Challenges Overcome During Development")

    challenges = [
        ("01", "Rate Limits on Free Cloud APIs",
         "PROBLEM: Google's free-tier embedding API frequently hit HTTP 429 ResourceExhausted limits during multi-page document ingestion.",
         "ENGINEERED FIX: Integrated local ONNX all-MiniLM-L6-v2 embeddings running directly on device. Result: Zero API quota, zero cost, and 15ms batch encoding speed."),
        ("02", "Reasoning-Model Output Loops in Ollama",
         "PROBLEM: Qwen 3.5 9B streamed its chain-of-thought into a thinking field for 90+ seconds, leaving the response field empty and triggering timeouts in LangChain.",
         "ENGINEERED FIX: Developed a custom OllamaDirectLLM client explicitly enforcing think=False. Response latency dropped from 90s to 7.5s with instant JSON output."),
        ("03", "Zero Hallucination with Verifiable Grounding",
         "PROBLEM: Standard LLM summarizers tend to invent deadlines or assume task owners not stated in the source text.",
         "ENGINEERED FIX: Designed a strict Pydantic extraction schema requiring an exact verbatim source quote for every action item, guaranteeing complete factual grounding.")
    ]

    for i, (num, c_title, c_prob, c_sol) in enumerate(challenges):
        card = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9 + i * 3.9), Inches(2.0), Inches(3.7), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tfc = card.text_frame
        tfc.margin_left = tfc.margin_top = tfc.margin_right = tfc.margin_bottom = Inches(0.25)
        tfc.word_wrap = True

        pn = tfc.paragraphs[0]
        pn.text = num
        pn.font.name = "Georgia"
        pn.font.size = Pt(24)
        pn.font.bold = True
        pn.font.color.rgb = ACCENT_CORAL

        pt = tfc.add_paragraph()
        pt.text = c_title
        pt.font.name = "Georgia"
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_TITLE
        pt.space_before = Pt(4)

        pp = tfc.add_paragraph()
        pp.text = c_prob
        pp.font.name = "Segoe UI"
        pp.font.size = Pt(11)
        pp.font.color.rgb = RGBColor(160, 50, 50)
        pp.space_before = Pt(8)

        ps = tfc.add_paragraph()
        ps.text = c_sol
        ps.font.name = "Segoe UI"
        ps.font.size = Pt(11)
        ps.font.color.rgb = RGBColor(30, 110, 50)
        ps.space_before = Pt(8)

    # =========================================================================
    # SLIDE 9: CONCLUSION & PROJECT SUMMARY
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9, use_prismatic=True)
    add_header(s9, "08 / Summary & Conclusion", "Document Action Extractor: Ready for Academic & Real-World Use")

    # Left Summary Card
    card_l = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.9), Inches(5.6), Inches(4.9))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = CARD_BG
    card_l.line.color.rgb = CARD_BORDER
    card_l.line.width = Pt(1)

    tfl9 = card_l.text_frame
    tfl9.margin_left = tfl9.margin_top = tfl9.margin_right = tfl9.margin_bottom = Inches(0.3)
    tfl9.word_wrap = True

    p = tfl9.paragraphs[0]
    p.text = "Key Project Achievements"
    p.font.name = "Georgia"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    achievements = [
        "Complete End-to-End RAG System: Seamless pipeline from raw file ingestion to vector retrieval and LLM extraction.",
        "100% Data Privacy: Full offline support ensuring corporate and confidential documents never leave the local machine.",
        "Zero-Hallucination Extraction: Strict source grounding ensures every single extracted task reflects verified document context.",
        "Dual Inference Engines: Flexible runtime choosing between local Ollama privacy and high-speed Google Gemini Flash.",
        "Clean, Modern Aesthetic: Dark slate visual design built for distraction-free user experience and clear decision making."
    ]

    for a in achievements:
        pa = tfl9.add_paragraph()
        pa.text = "✓ " + a
        pa.font.name = "Segoe UI"
        pa.font.size = Pt(11.5)
        pa.font.color.rgb = TEXT_BODY
        pa.space_before = Pt(8)

    # Right Presenters & Demo Card
    card_r = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.6), Inches(4.9))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = CARD_BG
    card_r.line.color.rgb = CARD_BORDER
    card_r.line.width = Pt(1)

    tfr9 = card_r.text_frame
    tfr9.margin_left = tfr9.margin_top = tfr9.margin_right = tfr9.margin_bottom = Inches(0.35)
    tfr9.word_wrap = True

    pr = tfr9.paragraphs[0]
    pr.text = "Project Team & Academic Submission"
    pr.font.name = "Georgia"
    pr.font.size = Pt(18)
    pr.font.bold = True
    pr.font.color.rgb = TEXT_TITLE

    p_names = tfr9.add_paragraph()
    p_names.text = "Presenters:\n• Keshav (Roll No: 2449391)\n• Darshveer (Roll No: 2449377)"
    p_names.font.name = "Segoe UI"
    p_names.font.size = Pt(14)
    p_names.font.bold = True
    p_names.font.color.rgb = ACCENT_CORAL
    p_names.space_before = Pt(12)

    p_dept = tfr9.add_paragraph()
    p_dept.text = "Department of Computer Science & Engineering\nCapstone Class Project Submission"
    p_dept.font.name = "Segoe UI"
    p_dept.font.size = Pt(12)
    p_dept.font.color.rgb = TEXT_MUTED
    p_dept.space_before = Pt(8)

    p_repo = tfr9.add_paragraph()
    p_repo.text = "GitHub Repository:\ngithub.com/keshav-x/rag-action-extractor"
    p_repo.font.name = "Segoe UI"
    p_repo.font.size = Pt(13)
    p_repo.font.bold = True
    p_repo.font.color.rgb = ACCENT_SLATE
    p_repo.space_before = Pt(14)

    p_q = tfr9.add_paragraph()
    p_q.text = "Thank you! Open for Questions & Live Demonstration."
    p_q.font.name = "Georgia"
    p_q.font.size = Pt(14)
    p_q.font.bold = True
    p_q.font.color.rgb = ACCENT_GREEN
    p_q.space_before = Pt(16)

    # Save presentations
    output_path = os.path.abspath("Document_Action_Extractor_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

    # Also save as a distinct copy
    copy_path = os.path.abspath("Document_Action_Extractor_Simplified.pptx")
    prs.save(copy_path)
    print(f"Second copy saved to: {copy_path}")

if __name__ == "__main__":
    create_presentation()
