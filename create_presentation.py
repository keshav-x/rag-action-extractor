import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette inspired by Prismatic Editorial theme
    BG_CREAM = RGBColor(250, 248, 245)
    TEXT_TITLE = RGBColor(30, 27, 38)        # Deep Plum / Slate
    TEXT_BODY = RGBColor(55, 53, 64)         # Charcoal Slate
    TEXT_MUTED = RGBColor(112, 108, 120)     # Muted Eyebrow
    CARD_BG = RGBColor(255, 255, 255)        # Pure White Card
    CARD_BORDER = RGBColor(226, 222, 215)    # Subtle Cream/Gray Border
    ACCENT_CORAL = RGBColor(215, 65, 85)     # Editorial Coral
    ACCENT_SLATE = RGBColor(45, 55, 72)

    bg_image_path = os.path.abspath("prismatic_bg.jpg")
    screenshot_main = os.path.abspath("slide_screenshot_main.png")
    screenshot_indexed = os.path.abspath("slide_screenshot_indexed.png")
    screenshot_results = os.path.abspath("slide_screenshot_results.png")

    def add_background(slide, use_prismatic=True):
        if use_prismatic and os.path.exists(bg_image_path):
            slide.shapes.add_picture(bg_image_path, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
        else:
            # Solid warm cream background
            bg_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
            bg_rect.fill.solid()
            bg_rect.fill.fore_color.rgb = BG_CREAM
            bg_rect.line.fill.background()

    def add_header(slide, eyebrow_text, title_text):
        # Eyebrow / Category
        tx_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = eyebrow_text.upper()
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CORAL

        # Title
        tx_title = slide.shapes.add_textbox(Inches(0.9), Inches(0.95), Inches(11.5), Inches(0.8))
        tf_title = tx_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_t = tf_title.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = "Georgia"
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_TITLE

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    add_background(s1, use_prismatic=True)

    # Subtitle Eyebrow
    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(7.5), Inches(0.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "ACADEMIC CAPSTONE PROJECT"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CORAL

    # Main Title
    tb_title = s1.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(8.5), Inches(2.2))
    tft = tb_title.text_frame
    tft.word_wrap = True
    p_main = tft.paragraphs[0]
    p_main.text = "Document Action\nExtractor"
    p_main.font.name = "Georgia"
    p_main.font.size = Pt(48)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_TITLE

    # Description
    tb_desc = s1.shapes.add_textbox(Inches(1.0), Inches(5.1), Inches(7.5), Inches(0.8))
    tfd = tb_desc.text_frame
    tfd.word_wrap = True
    pd = tfd.paragraphs[0]
    pd.text = "Automated Action Item, Assignee & Deadline Extraction\nPowered by Retrieval-Augmented Generation & Intelligent LLMs"
    pd.font.name = "Segoe UI"
    pd.font.size = Pt(15)
    pd.font.color.rgb = TEXT_BODY

    # Presenter Tag
    tb_pres = s1.shapes.add_textbox(Inches(1.0), Inches(6.3), Inches(7.5), Inches(0.5))
    tfp = tb_pres.text_frame
    pp = tfp.paragraphs[0]
    pp.text = "Presented by Keshav | Department of Computer Science & Engineering"
    pp.font.name = "Segoe UI"
    pp.font.size = Pt(11)
    pp.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: THE PROBLEM
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_background(s2, use_prismatic=False)
    add_header(s2, "01 / Background & Motivation", "The Challenge of Unstructured Business Documents")

    # Left Narrative Card
    card_l = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.9), Inches(5.6), Inches(5.0))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = CARD_BG
    card_l.line.color.rgb = CARD_BORDER
    card_l.line.width = Pt(1)

    tf_l = card_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = Inches(0.35)

    p = tf_l.paragraphs[0]
    p.text = "Information Overload in Team Meetings"
    p.font.name = "Georgia"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    points_l = [
        "Modern organizations produce hundreds of pages of meeting minutes, technical specifications, and sprint reviews every week.",
        "Crucial commitments, assigned deliverables, and firm deadlines remain buried inside multi-page paragraphs.",
        "Manual tracking is slow, expensive, and subject to human oversight, leading to forgotten tasks and delayed projects.",
        "Organizations lack an automated pipeline to instantly convert raw documents into clean, accountable action items."
    ]
    for pt in points_l:
        p = tf_l.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(12)

    # Right 3 Impact Cards
    impacts = [
        ("Loss of Task Ownership", "When tasks are mentioned colloquially, responsibility is frequently diffused across team members without explicit accountability."),
        ("Missed Deliverable Deadlines", "Timeframes ('by Friday', 'end of Q4') buried within narrative prose are frequently overlooked by project managers."),
        ("Inefficient Administrative Overhead", "Engineers and managers waste up to 4.5 hours per week manually extracting and cataloging action items.")
    ]
    for i, (imp_title, imp_desc) in enumerate(impacts):
        card_r = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.9 + i * 1.7), Inches(5.6), Inches(1.5))
        card_r.fill.solid()
        card_r.fill.fore_color.rgb = CARD_BG
        card_r.line.color.rgb = CARD_BORDER
        card_r.line.width = Pt(1)

        tfr = card_r.text_frame
        tfr.word_wrap = True
        tfr.margin_left = tfr.margin_top = tfr.margin_right = tfr.margin_bottom = Inches(0.2)

        pr = tfr.paragraphs[0]
        pr.text = f"0{i+1}. {imp_title}"
        pr.font.name = "Georgia"
        pr.font.size = Pt(15)
        pr.font.bold = True
        pr.font.color.rgb = ACCENT_CORAL

        pr2 = tfr.add_paragraph()
        pr2.text = imp_desc
        pr2.font.name = "Segoe UI"
        pr2.font.size = Pt(12)
        pr2.font.color.rgb = TEXT_BODY
        pr2.space_before = Pt(4)

    # ==========================================
    # SLIDE 3: THE SOLUTION
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_background(s3, use_prismatic=False)
    add_header(s3, "02 / The Proposed Solution", "RAG-Powered Autonomous Action Extraction")

    pillars = [
        ("Multi-Format Parsing", "Native ingestion supporting PDF, DOCX, and TXT documents with intelligent text extraction and semantic boundary preservation."),
        ("Vector Similarity Retrieval", "Recursively splits text into overlapping chunks, indexing into ChromaDB for high-precision semantic search against targeted queries."),
        ("Structured Schema Extraction", "Leverages advanced LLMs to extract strictly validated Pydantic models (Task, Owner, Deadline, Priority, and Source Citation)."),
        ("Auditable & Verifiable", "Every single extracted action item includes direct contextual evidence (exact sentence quote) guaranteeing zero hallucinations.")
    ]

    for i, (pil_title, pil_desc) in enumerate(pillars):
        x = Inches(0.9 + (i % 2) * 5.9)
        y = Inches(1.9 + (i // 2) * 2.5)

        card_p = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(5.6), Inches(2.2))
        card_p.fill.solid()
        card_p.fill.fore_color.rgb = CARD_BG
        card_p.line.color.rgb = CARD_BORDER
        card_p.line.width = Pt(1)

        tfp = card_p.text_frame
        tfp.word_wrap = True
        tfp.margin_left = tfp.margin_top = tfp.margin_right = tfp.margin_bottom = Inches(0.25)

        p = tfp.paragraphs[0]
        p.text = f"Pillar {i+1}: {pil_title}"
        p.font.name = "Georgia"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_TITLE

        p2 = tfp.add_paragraph()
        p2.text = pil_desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(8)

    # ==========================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_background(s4, use_prismatic=False)
    add_header(s4, "03 / Architecture & Workflow", "Comprehensive Technical Pipeline")

    steps = [
        ("Step 1: Ingestion", "Document Parsing", "Accepts .pdf (pypdf), .docx (python-docx), and .txt, extracting plain text stream."),
        ("Step 2: Chunking", "Text Splitting", "LangChain RecursiveCharacterTextSplitter splits into 1000-char chunks with 150-char overlap."),
        ("Step 3: Embedding", "Vector Indexing", "Paced batching (25 chunks/batch) stores embeddings in ChromaDB persistent client."),
        ("Step 4: Retrieval", "Similarity Search", "Top-K similarity retrieval pulls the most relevant context chunks for user queries."),
        ("Step 5: Reasoning", "LLM Inference", "Google Gemini Flash or Local Ollama Qwen executes structured JSON extraction."),
        ("Step 6: Delivery", "Validation & Export", "Pydantic ActionItemList validation, priority tagging, and 1-click CSV/JSON export.")
    ]

    for i, (s_num, s_title, s_desc) in enumerate(steps):
        col = i % 3
        row = i // 3
        x = Inches(0.9 + col * 3.9)
        y = Inches(1.9 + row * 2.5)

        c = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(3.65), Inches(2.2))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1)

        tfc = c.text_frame
        tfc.word_wrap = True
        tfc.margin_left = tfc.margin_top = tfc.margin_right = tfc.margin_bottom = Inches(0.2)

        p = tfc.paragraphs[0]
        p.text = s_num
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CORAL

        p2 = tfc.add_paragraph()
        p2.text = s_title
        p2.font.name = "Georgia"
        p2.font.size = Pt(15)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_TITLE
        p2.space_before = Pt(2)

        p3 = tfc.add_paragraph()
        p3.text = s_desc
        p3.font.name = "Segoe UI"
        p3.font.size = Pt(11.5)
        p3.font.color.rgb = TEXT_BODY
        p3.space_before = Pt(6)

    # ==========================================
    # SLIDE 5: TECHNICAL INNOVATIONS
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_background(s5, use_prismatic=False)
    add_header(s5, "04 / Engineering Breakthroughs", "Solving Rate Limits & Enabling Offline Execution")

    # Innovation 1
    c1 = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.9), Inches(5.6), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    c1.line.width = Pt(1)

    tfc1 = c1.text_frame
    tfc1.word_wrap = True
    tfc1.margin_left = tfc1.margin_top = tfc1.margin_right = tfc1.margin_bottom = Inches(0.35)

    p = tfc1.paragraphs[0]
    p.text = "Solving Google 429 Quota Exhaustion"
    p.font.name = "Georgia"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    p_body1 = [
        "The Problem: Google GenAI Free Tier imposes a strict 100 requests/minute quota. A typical 1MB PDF generates 200+ chunks, causing immediate 429 crashes.",
        "The Solution: Implemented local CPU embeddings using all-MiniLM-L6-v2 via ONNX.",
        "Zero API Calls: Entire document indexing runs 100% locally with 0 external API calls.",
        "Adaptive Backoff: For cloud embeddings, paced batches of 25 chunks with dynamic 40s+ retry handlers guarantee zero unhandled crashes."
    ]
    for b in p_body1:
        p = tfc1.add_paragraph()
        p.text = "• " + b
        p.font.name = "Segoe UI"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(10)

    # Innovation 2
    c2 = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.6), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER
    c2.line.width = Pt(1)

    tfc2 = c2.text_frame
    tfc2.word_wrap = True
    tfc2.margin_left = tfc2.margin_top = tfc2.margin_right = tfc2.margin_bottom = Inches(0.35)

    p = tfc2.paragraphs[0]
    p.text = "Dual-Engine Inference Architecture"
    p.font.name = "Georgia"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    p_body2 = [
        "Cloud LLM (Google Gemini): Lightning-fast inference (1.5s per query) utilizing Gemini 1.5 & 2.0 Flash for structured output generation.",
        "Local LLM (Ollama): Complete offline privacy with zero external connectivity requirements.",
        "Integrated Model: Connected to huihui_ai/qwen3.5-abliterated:9b running directly on local machine via Ollama API.",
        "100% Offline Capability: When paired with local embeddings, the application functions seamlessly even without an internet connection."
    ]
    for b in p_body2:
        p = tfc2.add_paragraph()
        p.text = "• " + b
        p.font.name = "Segoe UI"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(10)

    # ==========================================
    # SLIDE 6: PRODUCT DEMO - WORKSPACE & INGESTION
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_background(s6, use_prismatic=False)
    add_header(s6, "05 / Product Demonstration", "Minimalist User Interface & Document Ingestion")

    # Add Screenshot
    if os.path.exists(screenshot_indexed):
        pic = s6.shapes.add_picture(screenshot_indexed, Inches(0.9), Inches(1.9), width=Inches(7.2))

    # Right explanation card
    ce = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.4), Inches(1.9), Inches(4.0), Inches(5.0))
    ce.fill.solid()
    ce.fill.fore_color.rgb = CARD_BG
    ce.line.color.rgb = CARD_BORDER
    ce.line.width = Pt(1)

    tf_e = ce.text_frame
    tf_e.word_wrap = True
    tf_e.margin_left = tf_e.margin_top = tf_e.margin_right = tf_e.margin_bottom = Inches(0.3)

    p = tf_e.paragraphs[0]
    p.text = "Document Ingestion Workflow"
    p.font.name = "Georgia"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    demo_pts = [
        "1. Multi-Format Upload: Drag-and-drop support for PDF, Word (.docx), and plain text (.txt).",
        "2. Transparent Text Preview: Collapsible preview drawer lets users verify raw extracted text before querying.",
        "3. Live Status Indicators: Displays real-time chunk counts and vector database indexing state.",
        "4. One-Click Presets: Instant buttons for 'All Action Items', 'High Priority', and 'Upcoming Deadlines'."
    ]
    for dp in demo_pts:
        p = tf_e.add_paragraph()
        p.text = dp
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(12)

    # ==========================================
    # SLIDE 7: PRODUCT DEMO - RESULTS & CITATIONS
    # ==========================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_background(s7, use_prismatic=False)
    add_header(s7, "06 / Results & Citations", "Structured Action Items with Full Traceability")

    # Add Screenshot
    if os.path.exists(screenshot_results):
        pic2 = s7.shapes.add_picture(screenshot_results, Inches(0.9), Inches(1.9), width=Inches(7.2))

    # Right explanation card
    ce2 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.4), Inches(1.9), Inches(4.0), Inches(5.0))
    ce2.fill.solid()
    ce2.fill.fore_color.rgb = CARD_BG
    ce2.line.color.rgb = CARD_BORDER
    ce2.line.width = Pt(1)

    tf_e2 = ce2.text_frame
    tf_e2.word_wrap = True
    tf_e2.margin_left = tf_e2.margin_top = tf_e2.margin_right = tf_e2.margin_bottom = Inches(0.3)

    p = tf_e2.paragraphs[0]
    p.text = "Structured Output Features"
    p.font.name = "Georgia"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    res_pts = [
        "1. Complete Visibility: Full task descriptions wrap cleanly without horizontal scrollbars or cutoffs.",
        "2. Clear Ownership: Highlights exact assigned owners (e.g. Sarah Chen, Alex Rodriguez) and due dates.",
        "3. Priority Tags: High, Medium, and Low urgency badges for instantaneous task triage.",
        "4. Source Citations: Direct document quotations prevent LLM hallucinations and enable easy verification.",
        "5. 1-Click Export: Immediate download in structured CSV or JSON formats for project management tools."
    ]
    for rp in res_pts:
        p = tf_e2.add_paragraph()
        p.text = rp
        p.font.name = "Segoe UI"
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(10)

    # ==========================================
    # SLIDE 8: TECH STACK
    # ==========================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    add_background(s8, use_prismatic=False)
    add_header(s8, "07 / Technology Stack", "Modular & Scalable Modern Architecture")

    tech_items = [
        ("Streamlit", "Frontend Web Interface", "Lightweight, reactive Python framework delivering clean dark slate UI without web bloat."),
        ("LangChain", "RAG Orchestration", "Manages document chunking, prompt composition, vector query pipelines, and model adapters."),
        ("ChromaDB", "Vector Database", "Persistent local vector store storing document embeddings for fast semantic similarity search."),
        ("Google Gemini", "Cloud Reasoning", "State-of-the-art Flash models providing sub-2-second JSON structured extraction."),
        ("Ollama", "Local Offline LLM", "Runs open-weights Qwen 3.5 9B locally on CPU/GPU for complete data confidentiality."),
        ("Pydantic v2", "Schema Validation", "Strict type validation guaranteeing compliant JSON output without parsing failures.")
    ]

    for i, (t_name, t_role, t_detail) in enumerate(tech_items):
        col = i % 3
        row = i // 3
        x = Inches(0.9 + col * 3.9)
        y = Inches(1.9 + row * 2.5)

        ct = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(3.65), Inches(2.2))
        ct.fill.solid()
        ct.fill.fore_color.rgb = CARD_BG
        ct.line.color.rgb = CARD_BORDER
        ct.line.width = Pt(1)

        tfct = ct.text_frame
        tfct.word_wrap = True
        tfct.margin_left = tfct.margin_top = tfct.margin_right = tfct.margin_bottom = Inches(0.2)

        p = tfct.paragraphs[0]
        p.text = t_name
        p.font.name = "Georgia"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_TITLE

        p2 = tfct.add_paragraph()
        p2.text = t_role
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = ACCENT_CORAL
        p2.space_before = Pt(2)

        p3 = tfct.add_paragraph()
        p3.text = t_detail
        p3.font.name = "Segoe UI"
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_BODY
        p3.space_before = Pt(6)

    # ==========================================
    # SLIDE 9: CONCLUSION & FUTURE SCOPE
    # ==========================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    add_background(s9, use_prismatic=True)

    # Content Box
    cb = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.5), Inches(8.5), Inches(5.2))
    cb.fill.solid()
    cb.fill.fore_color.rgb = CARD_BG
    cb.line.color.rgb = CARD_BORDER
    cb.line.width = Pt(1)

    tfcb = cb.text_frame
    tfcb.word_wrap = True
    tfcb.margin_left = tfcb.margin_top = tfcb.margin_right = tfcb.margin_bottom = Inches(0.4)

    p = tfcb.paragraphs[0]
    p.text = "Summary & Future Roadmap"
    p.font.name = "Georgia"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    p_conc = [
        "Delivered a complete, functional RAG application extracting structured action items from unstructured business documents.",
        "Overcame Google free-tier 429 quota exhaustion through custom local embedding engine and adaptive request pacing.",
        "Engineered dual-engine inference enabling both cloud-scale Gemini processing and 100% offline Ollama execution.",
        "Future Roadmap: Integrations with Jira, Notion, and Slack APIs for direct automated task ticket creation.",
        "Thank you! Ready for Live Demonstration & Questions."
    ]
    for c in p_conc:
        p = tfcb.add_paragraph()
        p.text = "• " + c
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(12)

    output_file = "Document_Action_Extractor_Presentation.pptx"
    prs.save(output_file)
    print(f"Presentation saved successfully to: {os.path.abspath(output_file)}")

if __name__ == "__main__":
    build_presentation()
